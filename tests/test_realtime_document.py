from dataclasses import replace
import json
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from realtime import RealtimeEvent
from realtime_translation import TranslationEvent
from realtime_document import RealtimeDocument, DocumentStore
from docx import Document


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.events = []
        self.store = DocumentStore(self.directory / "recovery", lambda k, v: self.events.append((k, v)))
        self.store.new(True)
        self.final = RealtimeEvent("final_transcript", "core", 1, "Hello world.", 3601.25, 3603.5)
        self.translation = TranslationEvent("translation_completed", "core", 1, "Hello world.", "世界你好。",
                                            3601.25, 3603.5, "en", "zh-Hans")

    def tearDown(self):
        self.store.shutdown()
        self.temp.cleanup()

    def wait(self):
        deadline = time.monotonic() + 5
        while self.store.busy and time.monotonic() < deadline:
            time.sleep(0.005)
        self.assertFalse(self.store.busy)

    def test_creation_partial_exclusion_final_and_stop_retention(self):
        self.store.accept(replace(self.final, kind="partial_transcript"))
        self.assertEqual(self.store.snapshot().segments, [])
        self.store.accept(self.final)
        self.store.accept(self.final)
        self.store.finish()
        doc = self.store.snapshot()
        self.assertEqual(len(doc.segments), 1)
        self.assertIsNotNone(doc.ended_at)
        self.assertEqual(doc.segments[0].translation_status, "failed")
        self.assertEqual(doc.segments[0].end, 3603.5)

    def test_translation_alignment_and_early_translation(self):
        self.store.accept(self.translation)
        self.store.accept(self.final)
        segment = self.store.snapshot().segments[0]
        self.assertEqual(segment.translated_text, "世界你好。")
        self.assertEqual(segment.translation_status, "completed")
        with self.assertRaises(ValueError):
            self.store.accept(replace(self.translation, end=9000))

    def test_stale_session_does_not_modify_document(self):
        self.store.accept(self.final)
        self.store.accept(replace(self.translation, session_id="old"))
        self.assertTrue(self.store.snapshot().pending)

    def test_pending_export_requires_explicit_english_only(self):
        self.store.accept(self.final)
        with self.assertRaises(ValueError):
            self.store.snapshot().export(self.directory, ("docx",))
        paths = self.store.snapshot().export(self.directory, ("docx",), english_only=True)
        data = json.loads(next(p for p in paths if p.suffix == ".json").read_text(encoding="utf-8"))
        self.assertEqual(data["realtime_session"]["segments"][0]["translation_status"], "pending")
        self.assertIn("Hello", Document(next(p for p in paths if p.suffix == ".docx")).paragraphs[0].text)

    def test_all_formats_bilingual_docx_and_real_timestamps(self):
        self.store.accept(self.final)
        self.store.accept(self.translation)
        paths = self.store.snapshot().export(self.directory, ("txt", "docx", "json", "srt", "vtt"))
        by_format = {p.suffix: p for p in paths}
        self.assertIn("世界你好", by_format[".txt"].read_text(encoding="utf-8"))
        doc = Document(by_format[".docx"])
        self.assertEqual([p.text for p in doc.paragraphs], ["Hello world.", "世界你好。"])
        style = doc.styles["Normal"]
        self.assertEqual(style.font.name, "Calibri")
        self.assertEqual(style.font.size.pt, 11)
        self.assertEqual(style.paragraph_format.line_spacing, 1.15)
        self.assertIn("01:00:01,250 --> 01:00:03,500", by_format[".srt"].read_text(encoding="utf-8"))
        self.assertIn("01:00:01.250 --> 01:00:03.500", by_format[".vtt"].read_text(encoding="utf-8"))
        data = json.loads(by_format[".json"].read_text(encoding="utf-8"))
        self.assertEqual(RealtimeDocument.from_dict(data).segments[0].translated_text, "世界你好。")

    def test_failed_translation_preserves_readable_english(self):
        self.store.accept(self.final)
        self.store.accept(replace(self.translation, kind="translation_failed"))
        paths = self.store.snapshot().export(self.directory, ("txt", "docx"))
        text = paths[0].read_text(encoding="utf-8")
        self.assertIn("Hello", text)
        self.assertIn("翻译失败", text)

    def test_recovery_atomic_and_unicode(self):
        self.store.accept(self.final)
        self.store.accept(self.translation)
        self.wait()
        path = next((self.directory / "recovery").glob("*.json"))
        data = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(data["segments"][0]["translated_text"], "世界你好。")
        self.store.clear()
        self.store.restore(path)
        self.wait()
        self.assertEqual(len(self.store.snapshot().segments), 1)
        self.assertTrue(path.exists())
        self.assertEqual(list(path.parent.glob("*.tmp")), [])

    def test_new_session_preserves_old_recovery_and_unique_names(self):
        self.store.accept(self.final)
        old = self.store.snapshot()
        self.store.new()
        self.wait()
        self.assertTrue((self.directory / "recovery" / f"{old.session_id}.json").exists())
        first = old.export(self.directory, ("txt",), english_only=True)
        second = old.export(self.directory, ("txt",), english_only=True)
        self.assertNotEqual(first, second)
        self.assertTrue(all(p.exists() for p in first + second))

    def test_export_marks_only_saved_revision(self):
        self.store.accept(self.final)
        self.store.accept(self.translation)
        self.store.export(self.directory, ("txt",))
        self.wait()
        self.assertFalse(self.store.snapshot().unsaved)
        self.store.accept(replace(self.final, utterance_id=2, start=3604, end=3605))
        self.assertTrue(self.store.snapshot().unsaved)

    def test_corrupt_recovery_keeps_current_document(self):
        self.store.accept(self.final)
        path = self.directory / "bad.json"
        path.write_text("{}")
        self.store.restore(path)
        self.wait()
        self.assertEqual(len(self.store.snapshot().segments), 1)
        self.assertTrue(any(k == "error" for k, v in self.events))

    def test_failed_atomic_replace_preserves_previous_recovery(self):
        self.store.accept(self.final)
        self.wait()
        path = next((self.directory / "recovery").glob("*.json"))
        before = path.read_bytes()
        with patch("realtime_document.Path.replace", side_effect=PermissionError("locked")):
            self.store.accept(self.translation)
            self.wait()
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(self.store.snapshot().segments[0].translated_text, "世界你好。")
        self.store.autosave()
        self.wait()
        self.assertIsNone(self.store.io_error)


if __name__ == "__main__":
    unittest.main()
