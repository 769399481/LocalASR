"""No real model/audio needed. All file fixtures live in temporary directories."""

from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from batch import BatchProcessor
from docx import Document
from exporter import Exporter
from formatter import Formatter
from result import Segment, TranscriptionResult
from transcriber import Transcriber


def sample_result():
    text = "First sentence. Second sentence! Third sentence? Fourth sentence. Fifth sentence."
    return TranscriptionResult(text, "en", 8.0, (Segment(0.25, 8.0, text),))


class FormatterTests(unittest.TestCase):
    def test_preserves_tokens_and_punctuation(self):
        text = "example.com www.example.com https://example.com/a?b=c! user@example.com Dr.Smith U.S.A. 3.14 Hello!Why?"
        self.assertEqual(Formatter.clean(text), text)

    def test_natural_paragraphs(self):
        text = sample_result().text
        blocks = Formatter.clean(text).split("\n\n")
        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0], "First sentence. Second sentence! Third sentence? Fourth sentence.")
        self.assertEqual(" ".join(blocks), text)

    def test_abbreviations_do_not_count_as_sentences(self):
        text = "Dr. Smith met Mr. Jones in the U.S.A. today. It went well."
        self.assertEqual(Formatter.clean(text), text)

    def test_existing_paragraphs_and_idempotence(self):
        text = " First   paragraph\r\ncontinues.\r\n\r\nSecond paragraph. "
        cleaned = Formatter.clean(text)
        self.assertEqual(cleaned, "First paragraph continues.\n\nSecond paragraph.")
        self.assertEqual(Formatter.clean(cleaned), cleaned)

    def test_unpunctuated_long_text_is_bounded(self):
        text = "word " * 1000
        result = Formatter.clean(text)
        self.assertGreater(len(result.split("\n\n")), 1)
        self.assertEqual(result.split(), text.split())

    def test_empty(self):
        self.assertEqual(Formatter.clean(" \n\t"), "")


class FileTests(unittest.TestCase):
    @staticmethod
    def export_legacy(name, result, output):
        return Exporter.export(name, result, output, formats=("txt", "docx"))

    @staticmethod
    def is_complete_legacy(name, output):
        return Exporter.is_complete(name, output, formats=("txt", "docx"))

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="asr-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.input = self.root / "input"
        self.output = self.root / "output"
        self.input.mkdir()
        self.output.mkdir()

    def test_output_names(self):
        names = [Exporter.output_name(Path(p)) for p in
                 ("meeting.mp3", "meeting.wav", "meeting_mp3.wav")]
        self.assertEqual(names, ["meeting.mp3", "meeting.wav", "meeting_mp3.wav"])
        self.assertEqual(len(set(names)), 3)

    def test_completion_all_four_combinations(self):
        name = "meeting.mp3"
        self.assertFalse(self.is_complete_legacy(name, self.output))
        txt, docx = self.export_legacy(name, sample_result(), self.output)
        self.assertTrue(self.is_complete_legacy(name, self.output))
        saved = docx.read_bytes()
        docx.unlink()
        self.assertFalse(self.is_complete_legacy(name, self.output))
        docx.write_bytes(saved)
        txt.unlink()
        self.assertFalse(self.is_complete_legacy(name, self.output))

    def test_docx_format_and_utf8_roundtrip(self):
        result = sample_result()
        txt, docx = self.export_legacy("meeting.wav", result, self.output)
        document = Document(docx)
        self.assertEqual([p.text for p in document.paragraphs], Formatter.clean(result.text).split("\n\n"))
        self.assertEqual(txt.read_text(encoding="utf-8"), Formatter.clean(result.text))
        normal = document.styles["Normal"]
        self.assertEqual(normal.font.name, "Calibri")
        self.assertEqual(normal.font.size.pt, 11)
        self.assertAlmostEqual(normal.paragraph_format.line_spacing, 1.15)
        self.assertEqual(normal.paragraph_format.space_after.pt, 8)
        self.assertTrue(all(p.alignment == 0 for p in document.paragraphs))
        txt, docx = self.export_legacy("chinese.wav", "中文原文。", self.output)
        self.assertEqual(txt.read_text(encoding="utf-8"), "中文原文。")
        self.assertEqual(Document(docx).paragraphs[0].text, "中文原文。")

    def test_empty_transcription_is_valid(self):
        self.export_legacy("silence.wav", "", self.output)
        self.assertTrue(self.is_complete_legacy("silence.wav", self.output))

    def test_corrupt_or_mismatched_exports_are_incomplete(self):
        txt, docx = self.export_legacy("test", "original", self.output)
        txt.write_text("different", encoding="utf-8")
        self.assertFalse(self.is_complete_legacy("test", self.output))
        txt.write_text("original", encoding="utf-8")
        docx.write_bytes(b"broken")
        with self.assertLogs("exporter", level="WARNING"):
            self.assertFalse(self.is_complete_legacy("test", self.output))

    def test_staging_failure_preserves_existing_outputs(self):
        txt, docx = self.export_legacy("test", "old", self.output)
        original = (txt.read_bytes(), docx.read_bytes())
        with patch.object(Exporter, "save_docx", side_effect=OSError("disk failure")):
            with self.assertRaises(OSError):
                self.export_legacy("test", "new", self.output)
        self.assertEqual((txt.read_bytes(), docx.read_bytes()), original)
        self.assertFalse(list(self.output.glob(".asr-export-*")))

    def test_new_staging_failure_does_not_publish_txt(self):
        with patch.object(Exporter, "save_docx", side_effect=OSError("disk failure")):
            with self.assertRaises(OSError):
                self.export_legacy("test", "new", self.output)
        self.assertFalse(self.is_complete_legacy("test", self.output))
        self.assertEqual(list(self.output.iterdir()), [])

    def test_publication_failure_is_incomplete_and_recoverable(self):
        self.export_legacy("test", "old", self.output)
        original_replace = Path.replace

        def fail_docx(source, target):
            if source.name == "transcript.docx":
                raise PermissionError("Word has file open")
            return original_replace(source, target)

        with patch.object(Path, "replace", fail_docx):
            with self.assertRaises(PermissionError):
                self.export_legacy("test", "new", self.output)
        self.assertTrue((self.output / "test.pending").exists())
        self.assertFalse(self.is_complete_legacy("test", self.output))
        self.export_legacy("test", "new", self.output)
        self.assertTrue(self.is_complete_legacy("test", self.output))
        self.assertFalse((self.output / "test.pending").exists())

    def test_empty_batch_does_not_load_model(self):
        factory = Mock(side_effect=AssertionError("must not load"))
        summary = BatchProcessor(self.input, self.output, transcriber_factory=factory).run()
        self.assertEqual(summary.found, 0)
        factory.assert_not_called()

    def test_skip_does_not_load_model(self):
        (self.input / "test.mp3").touch()
        self.export_legacy("test.mp3", "done", self.output)
        factory = Mock(side_effect=AssertionError("must not load"))
        summary = BatchProcessor(self.input, self.output, transcriber_factory=factory,
                                 output_formats=("txt", "docx")).run()
        self.assertEqual((summary.found, summary.skipped, summary.processed), (1, 1, 0))
        factory.assert_not_called()

    def test_batch_success_failure_continuation_and_log_privacy(self):
        for name in ("bad.mp3", "meeting.mp3", "meeting.wav", "CAPS.WAV", "ignore.pdf"):
            (self.input / name).touch()
        (self.input / "directory.mp3").mkdir()
        engine = Mock()
        secret = "PRIVATE TRANSCRIPT CONTENT"

        def transcribe(path):
            if path.name == "bad.mp3":
                raise ValueError("invalid media")
            return TranscriptionResult(secret, "en", 1, (Segment(0, 1, secret),))

        engine.transcribe.side_effect = transcribe
        factory = Mock(return_value=engine)
        with self.assertLogs("batch", level="INFO") as captured:
            summary = BatchProcessor(self.input, self.output, transcriber_factory=factory).run()
        self.assertEqual((summary.found, summary.success, summary.failed, summary.processed), (4, 3, 1, 4))
        self.assertEqual(summary.failures, ["bad.mp3"])
        self.assertNotIn(secret, "\n".join(captured.output))
        self.assertIn("Traceback", "\n".join(captured.output))
        factory.assert_called_once()
        for name in ("meeting.mp3", "meeting.wav", "CAPS.WAV"):
            self.assertTrue(self.is_complete_legacy(name, self.output))

    def test_partial_output_is_reprocessed(self):
        (self.input / "test.wav").touch()
        (self.output / "test.wav.txt").write_text("old", encoding="utf-8")
        factory = Mock(return_value=Mock(transcribe=Mock(return_value=sample_result())))
        summary = BatchProcessor(self.input, self.output, transcriber_factory=factory).run()
        self.assertEqual((summary.success, summary.skipped), (1, 0))
        self.assertTrue(self.is_complete_legacy("test.wav", self.output))

    def test_model_initialization_failure_counts_all_pending(self):
        for name in ("a.wav", "b.wav"):
            (self.input / name).touch()
        with self.assertLogs("batch", level="ERROR"):
            summary = BatchProcessor(self.input, self.output,
                                     transcriber_factory=Mock(side_effect=RuntimeError("bad model"))).run()
        self.assertEqual((summary.found, summary.failed, summary.success), (2, 2, 0))
        self.assertEqual(summary.failures, ["a.wav", "b.wav"])

    def test_export_failure_is_reported(self):
        (self.input / "test.wav").touch()
        factory = Mock(return_value=Mock(transcribe=Mock(return_value=sample_result())))
        with patch.object(Exporter, "export", side_effect=PermissionError("locked")):
            with self.assertLogs("batch", level="ERROR"):
                summary = BatchProcessor(self.input, self.output, transcriber_factory=factory).run()
        self.assertEqual((summary.failed, summary.success), (1, 0))


class TranscriberTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="asr-model-test-")
        self.addCleanup(self.temporary.cleanup)
        self.model_path = Path(self.temporary.name)
        (self.model_path / "model.bin").touch()

    def make_engine(self, **kwargs):
        return Transcriber(model_path=self.model_path, **kwargs)

    def test_structured_segments_and_auto_language(self):
        with patch("transcriber.WhisperModel") as model:
            model.return_value.transcribe.return_value = (
                iter([SimpleNamespace(start=0.5, end=1.75, text=" Hello. "),
                      SimpleNamespace(start=2.0, end=3.0, text=" World. ")]),
                SimpleNamespace(language="en", duration=3.5))
            engine = self.make_engine(language="auto")
            result = engine.transcribe(Path("sample.wav"))
            self.assertEqual(result.text, "Hello.\nWorld.")
            self.assertEqual((result.language, result.duration), ("en", 3.5))
            self.assertEqual(result.segments[0], Segment(0.5, 1.75, "Hello."))
            self.assertIsNone(model.return_value.transcribe.call_args.kwargs["language"])
            self.assertTrue(model.call_args.kwargs["local_files_only"])

    def test_gpu_loading_falls_back_to_cpu(self):
        with patch("transcriber.WhisperModel", side_effect=[RuntimeError("CUDA unavailable"), Mock()]) as model:
            with self.assertLogs("transcriber", level="WARNING"):
                engine = self.make_engine(device="cuda", compute_type="float16")
            self.assertEqual(engine.device, "cpu")
            self.assertEqual(model.call_args.kwargs["compute_type"], "int8")

    def test_deferred_gpu_failure_falls_back(self):
        def broken_segments():
            yield SimpleNamespace(start=0, end=1, text="partial")
            raise RuntimeError("CUDA out of memory")

        gpu = Mock()
        gpu.transcribe.return_value = (broken_segments(), SimpleNamespace(language="en", duration=2))
        cpu = Mock()
        cpu.transcribe.return_value = (
            iter([SimpleNamespace(start=0, end=2, text="complete")]),
            SimpleNamespace(language="en", duration=2))
        with patch("transcriber.WhisperModel", side_effect=[gpu, cpu]):
            engine = self.make_engine(device="cuda", compute_type="float16")
            with self.assertLogs("transcriber", level="WARNING"):
                result = engine.transcribe("test.wav")
        self.assertEqual(result.text, "complete")
        self.assertEqual(engine.device, "cpu")

    def test_non_gpu_errors_propagate(self):
        with patch("transcriber.WhisperModel") as model:
            model.return_value.transcribe.side_effect = RuntimeError("invalid audio")
            engine = self.make_engine(device="cuda")
            with self.assertRaisesRegex(RuntimeError, "invalid audio"):
                engine.transcribe("bad.wav")
            self.assertEqual(model.call_count, 1)


if __name__ == "__main__":
    unittest.main()

