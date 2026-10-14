"""Translation tests use temporary output and injected providers, never a network."""

from dataclasses import replace
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from batch import BatchProcessor
from docx import Document
from exporter import Exporter
from result import Segment, TranscriptionResult
from translation import (MarianLocalTranslator, TranslationOptions, TranslationService,
                         create_translator)


def source(count=5):
    segments = tuple(Segment(i * 2 + 0.125, i * 2 + 1.875, f"Sentence {i + 1}.") for i in range(count))
    return TranscriptionResult("\n".join(s.text for s in segments), "en", count * 2, segments)


class TranslationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="asr-translation-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.output = self.root / "output"
        self.input = self.root / "input"
        self.input.mkdir()
        self.options = TranslationOptions(enabled=True, model_path=self.root / "model")
        self.request = self.options.completion_key()
        self.backend = Mock()
        self.backend.translate.side_effect = lambda text, src, dst: "中文译文：" + text
        self.factory = Mock(return_value=self.backend)
        self.original = source()

    def translate(self, original=None):
        return TranslationService(self.options, self.factory).translate(
            self.original if original is None else original, self.request)

    def batch(self, options=None, formats=("txt", "docx", "json", "srt", "vtt"), factory=None):
        (self.input / "test.wav").touch()
        return BatchProcessor(self.input, self.output,
                              transcriber_factory=Mock(return_value=Mock(transcribe=Mock(return_value=self.original))),
                              translation_options=self.options if options is None else options,
                              translator_factory=self.factory if factory is None else factory,
                              output_formats=formats)

    def test_segment_alignment_unicode_and_full_original(self):
        translated = self.translate()
        self.assertTrue(translated.complete)
        translated.validate_against(self.original)
        self.assertEqual(translated.original_text, self.original.text)
        self.assertEqual(translated.source_language, "en")
        self.assertEqual(translated.target_language, "zh-Hans")
        for old, new in zip(self.original.segments, translated.segments):
            self.assertEqual((old.start, old.end, old.text), (new.start, new.end, new.original_text))
            self.assertIn("中文译文", new.translated_text)
        self.assertEqual(self.original, source())

    def test_empty_segments_do_not_load_backend(self):
        for original in (TranscriptionResult("", "en", 0, ()),
                         TranscriptionResult(" \n", "en", 2, (Segment(0, 2, " \n"),))):
            with self.subTest(original=original):
                translated = self.translate(original)
                self.assertTrue(translated.complete)
                translated.validate_against(original)
        self.factory.assert_not_called()

    def test_missing_segments_with_nonempty_full_text_is_not_complete(self):
        translated = self.translate(TranscriptionResult("Do not lose this.", "en", 1, ()))
        self.assertFalse(translated.complete)
        self.assertEqual(translated.original_text, "Do not lose this.")
        self.factory.assert_not_called()

    def test_single_segment_failure_continues_and_logs_no_private_text(self):
        self.backend.translate.side_effect = ["第一段", ConnectionError("PRIVATE NETWORK BODY"), "第三段", "第四段", "第五段"]
        with self.assertLogs("translation", level="WARNING") as logs:
            translated = self.translate()
        self.assertFalse(translated.complete)
        self.assertIsNone(translated.segments[1].translated_text)
        self.assertEqual(translated.segments[2].translated_text, "第三段")
        self.assertIn("ConnectionError", translated.segments[1].error)
        self.assertNotIn("PRIVATE NETWORK BODY", "\n".join(logs.output))
        translated.validate_against(self.original)
        self.assertEqual(self.backend.translate.call_count, 5)

    def test_empty_backend_response_is_failed_not_success(self):
        self.backend.translate.return_value = ""
        self.backend.translate.side_effect = None
        with self.assertLogs("translation", level="WARNING"):
            translated = self.translate()
        self.assertFalse(translated.complete)
        self.assertTrue(all(seg.error is not None for seg in translated.segments))

    def test_unavailable_backend_preserves_original_and_loads_once(self):
        self.factory.side_effect = FileNotFoundError("model missing")
        service = TranslationService(self.options, self.factory)
        with self.assertLogs("translation", level="ERROR"):
            first = service.translate(self.original, self.request)
        second = service.translate(self.original, self.request)
        self.factory.assert_called_once()
        self.assertFalse(first.complete)
        self.assertEqual(first.original_text, second.original_text)
        first.validate_against(self.original)

    def test_source_language_mismatch_does_not_translate(self):
        original = replace(self.original, language="zh")
        translated = self.translate(original)
        self.assertEqual(translated.error, "source_language_mismatch")
        self.factory.assert_not_called()
        translated.validate_against(original)

    def test_unknown_backend_fails_explicitly(self):
        with self.assertRaises(ValueError):
            create_translator(replace(self.options, backend="not-installed"))
        with self.assertRaises(ValueError):
            create_translator(replace(self.options, target_language="ja"))

    def test_missing_local_model_does_not_download(self):
        with self.assertRaises(FileNotFoundError):
            MarianLocalTranslator(self.root / "missing-model")

    def test_json_additive_compatibility_and_english_subtitles(self):
        Exporter.export("test.wav", self.original, self.output)
        english = {fmt: (self.output / f"test.wav.{fmt}").read_bytes() for fmt in ("txt", "srt", "vtt")}
        data_before = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
        translated = self.translate()
        Exporter.export("test.wav", self.original, self.output, translation=translated)
        data = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
        self.assertEqual({key: data[key] for key in data_before}, data_before)
        self.assertEqual(data["translation"]["original_text"], self.original.text)
        self.assertTrue(data["translation"]["complete"])
        self.assertIn("中文译文", data["translation"]["translated_text"])
        for fmt, content in english.items():
            self.assertEqual((self.output / f"test.wav.{fmt}").read_bytes(), content)

    def test_bilingual_docx_groups_segments_into_paragraphs(self):
        translated = self.translate()
        Exporter.export("test.wav", self.original, self.output, translation=translated)
        document = Document(self.output / "test.wav.docx")
        texts = [p.text for p in document.paragraphs]
        self.assertEqual(len(texts), 4)  # 4 English segments grouped, then the last one.
        self.assertEqual(texts[0], "Sentence 1. Sentence 2. Sentence 3. Sentence 4.")
        self.assertIn("中文译文", texts[1])
        self.assertEqual(texts[2], "Sentence 5.")
        self.assertEqual(document.styles["Normal"].font.size.pt, 11)

    def test_export_rejects_misaligned_translation_before_writing(self):
        translated = self.translate()
        broken = replace(translated, segments=(replace(translated.segments[0], start=9),) + translated.segments[1:])
        with self.assertRaises(ValueError):
            Exporter.export("test.wav", self.original, self.output, translation=broken)
        self.assertFalse(self.output.exists())

    def test_completion_tracks_translation_options_and_disabled_state(self):
        Exporter.export("test.wav", self.original, self.output)
        self.assertFalse(Exporter.is_complete("test.wav", self.output, translation_request=self.request))
        translated = self.translate()
        Exporter.export("test.wav", self.original, self.output, translation=translated)
        self.assertTrue(Exporter.is_complete("test.wav", self.output, translation_request=self.request))
        self.assertFalse(Exporter.is_complete("test.wav", self.output))
        for key in ("target_language", "source_language", "backend", "revision", "model_path"):
            changed = {**self.request, key: "changed"}
            self.assertFalse(Exporter.is_complete("test.wav", self.output, translation_request=changed))

    def test_checkpoint_and_failed_translation_are_not_complete(self):
        Exporter.export("test.wav", self.original, self.output, translation_request=self.request)
        self.assertFalse(Exporter.is_complete("test.wav", self.output, translation_request=self.request))
        self.backend.translate.side_effect = RuntimeError("failed")
        with self.assertLogs("translation", level="WARNING"):
            translated = self.translate()
        Exporter.export("test.wav", self.original, self.output, translation=translated)
        self.assertFalse(Exporter.is_complete("test.wav", self.output, translation_request=self.request))
        self.assertIn("翻译失败", "\n".join(p.text for p in Document(self.output / "test.wav.docx").paragraphs))

    def test_disabled_batch_does_not_load_translator_or_change_json(self):
        processor = self.batch(options=replace(self.options, enabled=False))
        summary = processor.run()
        self.assertEqual((summary.success, summary.failed), (1, 0))
        self.factory.assert_not_called()
        data = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
        self.assertNotIn("translation", data)

    def test_enabled_batch_saves_english_before_translation_and_skips_when_complete(self):
        def translate(text, source_language, target_language):
            data = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
            self.assertEqual(data["text"], self.original.text)
            return "中文译文"

        self.backend.translate.side_effect = translate
        summary = self.batch().run()
        self.assertEqual((summary.success, summary.translation_failed), (1, 0))
        processor = self.batch(factory=Mock(side_effect=AssertionError("must not load")))
        processor.transcriber_factory = Mock(side_effect=AssertionError("must not load"))
        summary = processor.run()
        self.assertEqual(summary.skipped, 1)

    def test_failed_translation_preserves_asr_and_retries_next_run(self):
        self.factory.side_effect = RuntimeError("unavailable")
        with self.assertLogs(level="ERROR"):
            summary = self.batch().run()
        self.assertEqual((summary.failed, summary.success, summary.translation_failed), (1, 0, 1))
        data = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
        self.assertEqual(data["text"], self.original.text)
        self.assertFalse(data["translation"]["complete"])
        self.factory.side_effect = None
        summary = self.batch().run()
        self.assertEqual(summary.success, 1)

    def test_bilingual_export_failure_preserves_english_checkpoint(self):
        real_save = Exporter.save_json

        def fail_bilingual(result, path, **kwargs):
            if kwargs.get("translation") is not None:
                raise OSError("disk full")
            return real_save(result, path)

        with patch.object(Exporter, "save_json", side_effect=fail_bilingual), self.assertLogs("batch", level="ERROR"):
            summary = self.batch().run()
        self.assertEqual((summary.failed, summary.translation_failed), (1, 1))
        data = json.loads((self.output / "test.wav.json").read_text(encoding="utf-8"))
        self.assertEqual(data["text"], self.original.text)
        self.assertNotIn("translation", data)

    def test_translation_enabled_always_persists_json(self):
        summary = self.batch(formats=("srt",)).run()
        self.assertEqual(summary.success, 1)
        self.assertTrue((self.output / "test.wav.json").is_file())
        self.assertTrue((self.output / "test.wav.srt").is_file())

    def test_empty_batch_does_not_load_translator(self):
        processor = BatchProcessor(self.input, self.output, translation_options=self.options,
                                   translator_factory=self.factory)
        self.assertEqual(processor.run().found, 0)
        self.factory.assert_not_called()

    def test_long_local_segment_is_chunked_without_losing_input_tokens(self):
        import torch
        backend = MarianLocalTranslator.__new__(MarianLocalTranslator)
        backend.torch = torch
        backend.tokenizer = Mock()
        backend.tokenizer.encode.return_value = list(range(10, 610))
        backend.tokenizer.eos_token_id = 0
        backend.tokenizer.decode.return_value = "译文"
        backend.prefix_ids = [9]
        backend.max_positions = 512
        backend.model = Mock()
        backend.model.generate.return_value = torch.tensor([[0, 2, 0]])
        backend.simplify = Mock()
        backend.simplify.convert.side_effect = lambda text: text
        self.assertEqual(backend.translate("long segment", "en", "zh-Hans"), "译文译文译文")
        inputs = [call.kwargs["input_ids"][0].tolist()[1:-1] for call in backend.model.generate.call_args_list]
        self.assertEqual([token for chunk in inputs for token in chunk], list(range(10, 610)))
        backend.model.generate.return_value = torch.zeros((1, 512), dtype=torch.long)
        with self.assertRaises(RuntimeError):
            backend.translate("long segment", "en", "zh-Hans")


if __name__ == "__main__":
    unittest.main()
